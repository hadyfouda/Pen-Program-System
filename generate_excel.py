import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()

# Setup style
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
align_center = Alignment(horizontal="center", vertical="center")

table_style = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False,
)

# 1. Sheet: Enrolled_Accounts
ws1 = wb.active
ws1.title = "Enrolled_Accounts"
headers1 = [
    "account_id", "account_type", "account_name", "specialty", 
    "territory", "phone", "enrolled_by_rep", "pin_hash", 
    "allocated_pens", "dispensed_count", "current_balance", "status", "created_at"
]
ws1.append(headers1)
tab1 = Table(displayName="Tbl_Enrolled_Accounts", ref=f"A1:M2")
tab1.tableStyleInfo = table_style
# Add 1 empty placeholder row to make table valid
ws1.append(["", "", "", "", "", "", "", "", 0, 0, 0, "", ""])
ws1.add_table(tab1)

# 2. Sheet: Dispenses
ws2 = wb.create_sheet(title="Dispenses")
headers2 = [
    "record_id", "created_at", "account_id", "account_type", "account_name", 
    "recorded_by", "prescribing_doctor", "patient_name", "patient_phone", 
    "national_id", "product_name", "dispensed_qty", "rx_image_id", 
    "id_image_id", "consent_image_id", "status"
]
ws2.append(headers2)
tab2 = Table(displayName="Tbl_Dispenses", ref="A1:P2")
tab2.tableStyleInfo = table_style
ws2.append(["", "", "", "", "", "", "", "", "", "", "", 0, "", "", "", ""])
ws2.add_table(tab2)

# 3. Sheet: Stock_Requests
ws3 = wb.create_sheet(title="Stock_Requests")
headers3 = [
    "request_id", "created_at", "account_id", "account_name", "account_type", 
    "requested_by", "requested_qty", "approved_qty", "status", "reviewed_at"
]
ws3.append(headers3)
tab3 = Table(displayName="Tbl_Stock_Requests", ref="A1:J2")
tab3.tableStyleInfo = table_style
ws3.append(["", "", "", "", "", "", 0, 0, "", ""])
ws3.add_table(tab3)

# 4. Sheet: Exceptions
ws4 = wb.create_sheet(title="Exceptions")
headers4 = [
    "exception_id", "created_at", "account_id", "account_name", "patient_name", 
    "patient_phone", "national_id", "prescribing_doctor", "product_name", 
    "pens_count", "exception_reason", "admin_status", "admin_notes"
]
ws4.append(headers4)
tab4 = Table(displayName="Tbl_Exceptions", ref="A1:M2")
tab4.tableStyleInfo = table_style
ws4.append(["", "", "", "", "", "", "", "", "", 0, "", "", ""])
ws4.add_table(tab4)

# Set widths
for ws in [ws1, ws2, ws3, ws4]:
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

file_path = r"E:\Work\synapsevapharma\Pen-Program-System\Pen_Database.xlsx"
wb.save(file_path)
print("Pen_Database.xlsx created successfully at:", file_path)
